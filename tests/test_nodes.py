"""验证节点资产工作流与凭证关联节点启用测试。"""

from dataclasses import replace
from io import BytesIO
from pathlib import Path
import re
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from ngxops.accounts.models import User
from ngxops.accounts.passwords import make_password
from ngxops.app import create_app
from ngxops.config import get_settings
from ngxops.credentials.crypto import encrypt_secret
from ngxops.credentials.models import Credential
from ngxops.database.migration_runner import upgrade_database
from ngxops.database.session import session_scope
from ngxops.nodes.models import Node, NodeGroup
from ngxops.nodes.services import build_node_template_bytes
from ngxops.tasks.models import Task


@pytest.fixture
def node_client(tmp_path):
    """创建隔离数据库、超级管理员和已登录客户端。"""
    settings = replace(
        get_settings(),
        data_dir=tmp_path,
        resource_dir=Path.cwd(),
        debug=False,
        reload=False,
    )
    app = create_app(settings)
    with TestClient(app, follow_redirects=False) as client:
        upgrade_database(app.state.database.engine)
        with session_scope(app.state.database.session_factory) as session:
            with session.begin():
                user = User(
                    username="node-admin",
                    password=make_password("node-admin-password-123"),
                    is_superuser=True,
                    is_active=True,
                )
                session.add(user)
        login_page = client.get("/login/")
        token = re.search(r'name="csrf_token" value="([^"]+)"', login_page.text).group(1)
        login = client.post(
            "/login/",
            data={
                "csrf_token": token,
                "username": "node-admin",
                "password": "node-admin-password-123",
            },
        )
        assert login.status_code == 302
        yield client, app, token


def _node_form(csrf_token, **overrides):
    """生成节点新增和编辑共用的表单数据。"""
    values = {
        "csrf_token": csrf_token,
        "hostname": "node-a",
        "ip": "192.0.2.10",
        "port": "22",
        "environment": "prod",
        "nginx_path": "/usr/sbin/nginx",
        "main_conf_path": "/etc/nginx/nginx.conf",
        "credential_id": "",
        "description": "primary node",
    }
    values.update(overrides)
    return values


def test_node_crud_restore_groups_and_lock_gate(node_client):
    """验证节点创建、节点组成员、锁定、软删除和同 IP 恢复。"""
    client, app, csrf_token = node_client
    create = client.post("/nodes/create/", data=_node_form(csrf_token))
    assert create.status_code == 303, create.text
    with session_scope(app.state.database.session_factory) as session:
        node = session.scalar(select(Node).where(Node.ip == "192.0.2.10"))
        node_id = node.id
        user_id = node.created_by
        assert node.is_deleted is False
        assert node.status == "unknown"

    group_create = client.post(
        "/nodes/groups/create/",
        data={"csrf_token": csrf_token, "name": "production", "node_ids": [str(node_id)]},
    )
    assert group_create.status_code == 303, group_create.text
    group_list = client.get("/api/nodes/groups")
    assert group_list.status_code == 200
    assert group_list.json()["data"][0]["node_count"] == 1

    locked = client.post(
        "/api/nodes/lock",
        headers={"X-CSRFToken": csrf_token},
        json={"action": "lock", "node_ids": [node_id]},
    )
    assert locked.status_code == 200, locked.text
    with session_scope(app.state.database.session_factory) as session:
        node = session.get(Node, node_id)
        assert node.is_locked is True
        assert node.status == "offline"

    deleted = client.post(
        "/api/nodes/batch-delete",
        headers={"X-CSRFToken": csrf_token},
        json={"node_ids": [node_id]},
    )
    assert deleted.status_code == 200, deleted.text
    recreated = client.post(
        "/nodes/create/",
        data=_node_form(csrf_token, hostname="node-a-restored", description="restored"),
    )
    assert recreated.status_code == 303, recreated.text
    with session_scope(app.state.database.session_factory) as session:
        restored = session.get(Node, node_id)
        assert restored is not None
        assert restored.is_deleted is False
        assert restored.hostname == "node-a-restored"
        assert restored.created_by == user_id


def test_node_excel_template_export_and_atomic_import(node_client):
    """验证节点工作簿模板、导出和整份文件导入校验。"""
    client, app, csrf_token = node_client
    template = client.get("/api/nodes/import-template")
    assert template.status_code == 200
    assert "spreadsheetml" in template.headers["content-type"]

    workbook = load_workbook(BytesIO(build_node_template_bytes()))
    sheet = workbook["节点导入"]
    sheet.append(
        [
            "node-imported",
            "192.0.2.20",
            22,
            "生产",
            "/usr/sbin/nginx",
            "/etc/nginx/nginx.conf",
            "-",
            "-",
            "imported through template",
        ]
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    imported = client.post(
        "/api/nodes/import",
        headers={"X-CSRFToken": csrf_token},
        files={"file": ("nodes.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["created"] == 1
    exported = client.get("/api/nodes/export?search=node-imported")
    assert exported.status_code == 200
    exported_book = load_workbook(BytesIO(exported.content), data_only=True)
    exported_row = list(exported_book.active.iter_rows(values_only=True))[1]
    assert exported_row[0] == "node-imported"
    assert exported_row[7] in (None, "")

    duplicate_book = load_workbook(BytesIO(output.getvalue()))
    duplicate_book["节点导入"].append(
        ["node-duplicate", "192.0.2.20", 22, "测试", "", "", "", "", ""]
    )
    duplicate_output = BytesIO()
    duplicate_book.save(duplicate_output)
    duplicate_book.close()
    rejected = client.post(
        "/api/nodes/import",
        headers={"X-CSRFToken": csrf_token},
        files={"file": ("duplicate.xlsx", duplicate_output.getvalue())},
    )
    assert rejected.status_code == 422
    assert rejected.json()["success"] is False
    with session_scope(app.state.database.session_factory) as session:
        assert session.scalar(select(Node).where(Node.hostname == "node-duplicate")) is None


def test_credential_enable_tests_unlocked_nodes_and_updates_status(node_client, monkeypatch):
    """验证凭证启用任务进度、锁定跳过和节点双维度状态回写。"""
    client, app, csrf_token = node_client
    with session_scope(app.state.database.session_factory) as session:
        admin = session.scalar(select(User).where(User.username == "node-admin"))
        admin_id = admin.id
    with session_scope(app.state.database.session_factory) as session:
        with session.begin():
            credential = Credential(
                name="node-test-credential",
                username="deploy",
                auth_type="password",
                password=encrypt_secret(app.state.credential_encryption_key, "not-logged-password"),
                is_enabled=False,
                created_by=admin_id,
            )
            session.add(credential)
            session.flush()
            credential_id = credential.id
            session.add_all(
                [
                    Node(
                        hostname="testable-node",
                        ip="192.0.2.31",
                        port=22,
                        credential_id=credential_id,
                        environment="test",
                        created_by=admin_id,
                    ),
                    Node(
                        hostname="locked-node",
                        ip="192.0.2.32",
                        port=22,
                        credential_id=credential_id,
                        environment="test",
                        is_locked=True,
                        status="offline",
                        created_by=admin_id,
                    ),
                ]
            )

    class FakeClient:
        """提供不会建立远程连接的测试 SSH 客户端。"""

        def close(self):
            """模拟关闭 SSH 客户端。"""

    monkeypatch.setattr(
        "ngxops.credentials.tasks._connect_ssh",
        lambda *args: (FakeClient(), ""),
    )
    monkeypatch.setattr(
        "ngxops.credentials.tasks._probe_nginx",
        lambda client, path: (True, "1.26.1"),
    )
    enabled = client.post(
        "/api/credentials/{}/toggle-enable".format(credential_id),
        headers={"X-CSRFToken": csrf_token},
    )
    assert enabled.status_code == 200, enabled.text
    task_id = enabled.json()["task_id"]
    assert task_id
    progress = None
    for _ in range(60):
        progress_response = client.get(
            "/api/credentials/{}/enable-progress".format(credential_id)
        )
        assert progress_response.status_code == 200, progress_response.text
        progress = progress_response.json()
        if progress.get("status") not in ("pending", "running"):
            break
        time.sleep(0.05)
    with session_scope(app.state.database.session_factory) as session:
        task = session.get(Task, task_id)
        target_configs = task.target_configs
    assert progress["task_id"] == task_id, (progress, target_configs)
    assert progress["status"] == "success"
    assert progress["result_tree"]["summary"]["skipped"] == 1
    assert "not-logged-password" not in progress_response.text

    with session_scope(app.state.database.session_factory) as session:
        credential = session.get(Credential, credential_id)
        connected = session.scalar(select(Node).where(Node.ip == "192.0.2.31"))
        skipped = session.scalar(select(Node).where(Node.ip == "192.0.2.32"))
        assert credential.last_test_result == "success"
        assert connected.status == "online"
        assert connected.nginx_available is True
        assert connected.nginx_version == "1.26.1"
        assert skipped.status == "offline"

    disabled = client.post(
        "/api/credentials/{}/toggle-enable".format(credential_id),
        headers={"X-CSRFToken": csrf_token},
    )
    assert disabled.status_code == 200
    with session_scope(app.state.database.session_factory) as session:
        statuses = session.scalars(
            select(Node.status).where(Node.credential_id == credential_id)
        ).all()
        assert sorted(statuses) == ["offline", "offline"]
