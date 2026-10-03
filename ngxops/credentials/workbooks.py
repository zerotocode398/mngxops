"""提供凭证 Excel 模板、导入校验和受控明文导出。"""

import io
from typing import Any, Dict, List, Sequence, Tuple

import paramiko
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from ngxops.credentials.crypto import decrypt_secret, encrypt_secret
from ngxops.credentials.models import Credential


WORKBOOK_HEADERS = ("名称", "SSH用户", "认证方式", "密码", "私钥", "是否启用", "描述")
MAX_WORKBOOK_SIZE = 8 * 1024 * 1024
MAX_PASSWORD_LENGTH = 4096
MAX_PRIVATE_KEY_LENGTH = 65536
MAX_DESCRIPTION_LENGTH = 4000
_EMPTY_MARKERS = frozenset(("", "-", "—", "－", "无", "n/a", "na", "none"))


def build_credential_template_bytes() -> bytes:
    """生成包含固定表头和字段说明的凭证导入工作簿。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "凭证导入"
    sheet.append(WORKBOOK_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:G1"
    guide = workbook.create_sheet("填写说明")
    guide.append(["凭证批量导入说明"])
    guide.append(["表头和列顺序必须保持不变。"])
    guide.append(["名称、SSH用户、认证方式为必填；名称与当前用户已有凭证重名时会更新。"])
    guide.append(["认证方式填写密码认证/password 或密钥认证/key。"])
    guide.append(["密码认证填写密码，密钥认证填写未加密的 RSA、ECDSA 或 Ed25519 私钥。"])
    guide.append(["是否启用填写是/否；留空或填写 - 时默认启用。"])
    guide.append(["任一行校验失败时，整份文件不会导入；请妥善保管包含凭证明文的工作簿。"])
    guide.column_dimensions["A"].width = 100
    for sheet_item in (sheet, guide):
        for cell in sheet_item[1]:
            cell.font = Font(bold=True)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse_credential_workbook(content: bytes) -> Tuple[List[Dict[str, Any]], List[dict]]:
    """校验工作簿结构并提取凭证行且不在错误中回显单元格内容。"""
    if len(content) > MAX_WORKBOOK_SIZE:
        return [], [{"row": 0, "message": "文件不能超过 8 MiB"}]
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], [{"row": 0, "message": "无法读取 Excel 文件，请上传有效的 .xlsx"}]
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        normalized_header = tuple(_cell_text(value) for value in (header or ())[:7])
        if normalized_header != WORKBOOK_HEADERS:
            return [], [{"row": 1, "message": "表头不匹配，请使用凭证导入模板"}]
        parsed = []
        for row_number, raw in enumerate(rows, start=2):
            cells = list(raw or ())[:7]
            cells.extend([None] * (7 - len(cells)))
            if all(_is_empty(_cell_text(value)) for value in cells):
                continue
            parsed.append(
                {
                    "row": row_number,
                    "name": _cell_text(cells[0]),
                    "username": _cell_text(cells[1]),
                    "auth_type": _cell_text(cells[2]),
                    "password": _credential_cell_text(cells[3]).strip("\r\n"),
                    "private_key": _credential_cell_text(cells[4]).strip(),
                    "enabled": _cell_text(cells[5]),
                    "description": _cell_text(cells[6]),
                }
            )
        if not parsed:
            return [], [{"row": 0, "message": "没有可导入的数据行"}]
        return parsed, []
    except (StopIteration, KeyError, ValueError):
        return [], [{"row": 0, "message": "无法读取 Excel 工作表"}]
    finally:
        workbook.close()


def validate_credential_rows(
    rows: Sequence[Dict[str, Any]],
) -> Tuple[List[dict], List[dict]]:
    """整份校验凭证字段、同名项和认证材料后返回可导入数据。"""
    errors = []
    cleaned = []
    seen_names = set()
    for row in rows:
        row_errors = []
        name = str(row.get("name") or "").strip()
        username = str(row.get("username") or "").strip()
        auth_type = _normalize_auth_type(row.get("auth_type") or "")
        password = str(row.get("password") or "")
        private_key = str(row.get("private_key") or "").strip()
        description = str(row.get("description") or "").strip()
        enabled, enabled_valid = _normalize_enabled(row.get("enabled") or "")

        if not name:
            row_errors.append("名称不能为空")
        elif len(name) > 100:
            row_errors.append("名称长度不能超过 100 个字符")
        elif name in seen_names:
            row_errors.append("文件中存在重复名称")
        else:
            seen_names.add(name)
        if not username:
            row_errors.append("SSH 用户不能为空")
        elif len(username) > 100:
            row_errors.append("SSH 用户长度不能超过 100 个字符")
        if auth_type is None:
            row_errors.append("认证方式无效")
        if not enabled_valid:
            row_errors.append("是否启用请填写是或否")
        if len(password) > MAX_PASSWORD_LENGTH:
            row_errors.append("密码长度不能超过 4096 个字符")
        if len(private_key) > MAX_PRIVATE_KEY_LENGTH:
            row_errors.append("私钥内容不能超过 65536 个字符")
        if len(description) > MAX_DESCRIPTION_LENGTH:
            row_errors.append("描述不能超过 4000 个字符")
        if auth_type == "password":
            if not password.strip():
                row_errors.append("密码认证方式必须填写密码")
            private_key = ""
        elif auth_type == "key":
            if not private_key:
                row_errors.append("密钥认证方式必须填写私钥")
            elif not _is_valid_private_key(private_key):
                row_errors.append("私钥格式无效或包含口令")
            password = ""
        if row_errors:
            errors.extend(
                {"row": row["row"], "message": message} for message in row_errors
            )
            continue
        cleaned.append(
            {
                "row": row["row"],
                "name": name,
                "username": username,
                "auth_type": auth_type,
                "password": password,
                "private_key": private_key,
                "is_enabled": enabled,
                "description": description,
            }
        )
    return ([], errors) if errors else (cleaned, [])


def apply_credential_rows(
    session: Session,
    rows: Sequence[dict],
    owner_id: int,
    encryption_key: bytes,
) -> dict:
    """在调用方事务中更新当前用户同名凭证或创建新凭证。"""
    created_names = []
    updated_names = []
    for row in rows:
        existing = session.scalar(
            select(Credential).where(
                Credential.name == row["name"],
                Credential.created_by == owner_id,
            )
        )
        if existing is None:
            existing = Credential(name=row["name"], created_by=owner_id)
            session.add(existing)
            created_names.append(row["name"])
        else:
            updated_names.append(row["name"])
        existing.username = row["username"]
        existing.auth_type = row["auth_type"]
        existing.password = encrypt_secret(encryption_key, row["password"])
        existing.private_key = encrypt_secret(encryption_key, row["private_key"])
        existing.is_enabled = row["is_enabled"]
        existing.description = row["description"]
    return {
        "created": len(created_names),
        "updated": len(updated_names),
        "total": len(created_names) + len(updated_names),
        "created_names": created_names,
        "updated_names": updated_names,
    }


def build_credential_export_bytes(
    credentials: Sequence[Credential], encryption_key: bytes
) -> bytes:
    """生成仅供超级管理员下载的凭证明文 xlsx 内容。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "凭证导出"
    sheet.append(WORKBOOK_HEADERS)
    for credential in credentials:
        password = (
            decrypt_secret(encryption_key, credential.password)
            if credential.auth_type == "password"
            else ""
        )
        private_key = (
            decrypt_secret(encryption_key, credential.private_key)
            if credential.auth_type == "key"
            else ""
        )
        sheet.append(
            [
                credential.name,
                credential.username,
                "密码认证" if credential.auth_type == "password" else "密钥认证",
                password,
                private_key,
                "是" if credential.is_enabled else "否",
                credential.description,
            ]
        )
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def credential_name_preview(names: Sequence[str], limit: int = 20) -> str:
    """生成不包含认证材料的凭证名称预览。"""
    preview = "、".join(name for name in names[:limit] if name)
    if len(names) > limit:
        preview = "{} 等 {} 个".format(preview, len(names)) if preview else "共 {} 个".format(len(names))
    return preview


def _cell_text(value: Any) -> str:
    """将 Excel 单元格转成文本并清理首尾空白。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _credential_cell_text(value: Any) -> str:
    """转换认证材料单元格并保留密码中的有效空格。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _is_empty(value: str) -> bool:
    """识别工作簿中的空值及常用占位符。"""
    return (value or "").strip().lower() in _EMPTY_MARKERS


def _normalize_auth_type(value: str):
    """将表格中的认证方式名称归一化为内部枚举。"""
    normalized = (value or "").strip().lower()
    if normalized in ("password", "密码", "密码认证"):
        return "password"
    if normalized in ("key", "密钥", "秘钥", "密钥认证", "秘钥认证", "私钥"):
        return "key"
    return None


def _normalize_enabled(value: str) -> Tuple[bool, bool]:
    """将是否启用文本解析为布尔值并说明其有效性。"""
    normalized = (value or "").strip().lower()
    if normalized in ("", "-"):
        return True, True
    if normalized in ("是", "启用", "已启用", "true", "1", "yes", "y", "on"):
        return True, True
    if normalized in ("否", "禁用", "已禁用", "false", "0", "no", "n", "off"):
        return False, True
    return False, False


def _is_valid_private_key(value: str) -> bool:
    """验证未加密的 RSA、ECDSA 或 Ed25519 私钥。"""
    from io import StringIO

    for key_type in (paramiko.RSAKey, paramiko.ECDSAKey, paramiko.Ed25519Key):
        try:
            key_type.from_private_key(StringIO(value))
            return True
        except Exception:
            continue
    return False
